"""Assemble Appendix A (English layer of week3/APPENDIX_A_draft.md) into a Word file with Word equations.
Displays \\[ ... \\] become MTDisplayEquation paragraphs; tagged ones carry the MathType-style number field.
Output: revision_2026-09-26_ijpr/week3/Appendix_A_2026-09-27.docx (new file; no existing Word file is touched)."""
import copy
import re
import sys
from pathlib import Path

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).parent))
from s4_math import omml  # noqa: E402

SRC = Path(r"F:/Paper 3/revision_2026-09-26_ijpr/week3/APPENDIX_A_draft.md")
OUT = Path(r"F:/Paper 3/revision_2026-09-26_ijpr/week3/Appendix_A_2026-09-27.docx")
FONT, SIZE = "Times New Roman", 10.5
TEST_ONLY = "--test" in sys.argv

doc = docx.Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(3.0)
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(SIZE)
normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
st = doc.styles.add_style("MTDisplayEquation", WD_STYLE_TYPE.PARAGRAPH)
st._element.get_or_add_pPr().append(parse_xml(
    '<w:tabs xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:tab w:val="center" w:pos="4240"/><w:tab w:val="right" w:pos="8500"/></w:tabs>'))

INLINE = re.compile(r"(\\\(.*?\\\)|\*\*.*?\*\*|\*[^*\s][^*]*?\*)")
FIELD = '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:fldChar w:fldCharType="{t}"/></w:r>'
INSTR = '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:instrText xml:space="preserve">{x}</w:instrText></w:r>'
stats = {"inline": 0, "display": 0, "tagged": [], "table": 0, "heading": 0, "para": 0}


def set_font(run, size):
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    run.font.size = Pt(size)


M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WIDE = {"A.5"}                                   # displays too wide for the centre tab: set flush left
TABLE_COLS_CM = [1.75, 1.6, 2.5, 1.0, 1.6, 1.05, 1.1, 1.2, 1.8, 1.4]
CELL_MARGIN_TWIPS = 57                           # 0.1 cm left and right


def amath(tex, display, size=None):
    """MML2OMML pairs bare '|' operators into delimiters of mixed height; \\vert keeps them as plain bars."""
    o = copy.deepcopy(omml(re.sub(r"(?<!\\)\|", r"\\vert ", tex), display=display))
    if size:
        for r in o.iter(f"{{{M_NS}}}r"):
            rpr = r.find(f"{{{W_NS}}}rPr")
            if rpr is None:
                rpr = r.makeelement(f"{{{W_NS}}}rPr", {})
                mrpr = r.find(f"{{{M_NS}}}rPr")
                r.insert(0 if mrpr is None else 1, rpr)
            for tag in ("sz", "szCs"):
                e = rpr.makeelement(f"{{{W_NS}}}{tag}", {f"{{{W_NS}}}val": str(int(size * 2))})
                rpr.append(e)
    return o


def add_inline(par, text, size=SIZE, bold=False, italic=False):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("\\(") and tok.endswith("\\)"):
            par._p.append(amath(tok[2:-2], False, size if size != SIZE else None))
            stats["inline"] += 1
        elif tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            add_inline(par, tok[2:-2], size, True, italic)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            add_inline(par, tok[1:-1], size, bold, True)
        else:
            r = par.add_run(tok.replace("'", "\u2019"))
            set_font(r, size)
            r.bold = bold or None
            r.italic = italic or None


def display(tex):
    m = re.search(r"\\tag\{([^}]*)\}", tex)
    p = doc.add_paragraph(style="MTDisplayEquation")
    if not (m and m.group(1) in WIDE):
        p.add_run().add_tab()
    p._p.append(amath(tex, True))
    if m:
        p.add_run().add_tab()
        for xml in (FIELD.format(t="begin"), INSTR.format(x=" MACROBUTTON MTPlaceRef \\* MERGEFORMAT "),
                    FIELD.format(t="begin"), INSTR.format(x=" SEQ MTEqn \\h \\* MERGEFORMAT "),
                    FIELD.format(t="end"), INSTR.format(x=f"({m.group(1)})"), FIELD.format(t="end")):
            p._p.append(parse_xml(xml))
        stats["tagged"].append(m.group(1))
    stats["display"] += 1


def table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows if not re.match(r"^\|\s*-", r)]
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = "Table Grid"
    t.autofit = False
    tblpr = t._tbl.tblPr
    look = tblpr.find(qn("w:tblLook"))                  # schema order: ... tblLayout (set by autofit), tblCellMar, tblLook
    for xml in (f'<w:tblCellMar xmlns:w="{W_NS}"><w:left w:w="{CELL_MARGIN_TWIPS}" w:type="dxa"/>'
                f'<w:right w:w="{CELL_MARGIN_TWIPS}" w:type="dxa"/></w:tblCellMar>',):
        if look is None:
            tblpr.append(parse_xml(xml))
        else:
            look.addprevious(parse_xml(xml))
    widths = TABLE_COLS_CM if len(cells[0]) == len(TABLE_COLS_CM) else [15.0 / len(cells[0])] * len(cells[0])
    for i, row in enumerate(cells):
        trpr = t.rows[i]._tr.get_or_add_trPr()
        trpr.append(parse_xml(f'<w:cantSplit xmlns:w="{W_NS}"/>'))
        if i == 0:
            trpr.append(parse_xml(f'<w:tblHeader xmlns:w="{W_NS}"/>'))
        for j, txt in enumerate(row):
            cell = t.cell(i, j)
            cell.width = Cm(widths[j])
            par = cell.paragraphs[0]
            par.paragraph_format.keep_with_next = i < len(cells) - 1      # keep the small table on one page
            add_inline(par, txt, size=8, bold=(i == 0))
    doc.add_paragraph()
    stats["table"] += 1


def heading(text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    add_inline(p, text, size={1: 12, 2: 11, 3: 10.5}[level], bold=True)
    stats["heading"] += 1


def para(text, indent_first=True, left=0.0, prefix=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if left:
        p.paragraph_format.left_indent = Cm(left)
    add_inline(p, (("    " if indent_first and not prefix else "") + prefix + text))
    stats["para"] += 1


s = SRC.read_text(encoding="utf-8")
s = s[s.index("\n---\n", 4) + 5: s.index("## Review notes")]
lines = s.splitlines()
i = 0
prev = "heading"
while i < len(lines):
    ln = lines[i]
    if not ln.strip() or ln.lstrip().startswith(">") or ln.strip() == "---":
        i += 1
        continue
    if ln.strip().startswith("\\["):
        j = i
        while not lines[j].strip().endswith("\\]"):
            j += 1
        tex = " ".join(lines[i:j + 1]).strip()[2:-2]
        display(tex)
        prev = "display"
        i = j + 1
        continue
    if ln.startswith("#"):
        level = len(ln) - len(ln.lstrip("#"))
        heading(ln.lstrip("#").strip().split(" / ")[0], level)
        prev = "heading"
        i += 1
        continue
    if ln.startswith("|"):
        j = i
        while j < len(lines) and lines[j].startswith("|"):
            j += 1
        table(lines[i:j])
        prev = "table"
        i = j
        continue
    m = re.match(r"^(\s*)(- |\d+\. )(.*)$", ln)
    if m:
        bullet = m.group(2) == "- "
        para(m.group(3), left=0.6 + 0.4 * (len(m.group(1)) // 2), prefix=("\u2022 " if bullet else m.group(2)))
        prev = "list"
        i += 1
        continue
    text = ln.strip()
    para(text, indent_first=(prev != "display" and not text.startswith(("**", "*Proof", "("))))
    prev = "para"
    i += 1

if not TEST_ONLY:
    doc.save(OUT)
print(("[test] " if TEST_ONLY else f"wrote {OUT}\n") + str(stats))
